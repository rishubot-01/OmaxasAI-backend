
"""Extract text from files and text sources for OmaxasAI Knowledge."""

import io
import logging
from pathlib import Path

from apps.knowledge.models import KnowledgeDocument

logger = logging.getLogger(__name__)

MAX_DOCUMENT_SIZE = 20 * 1024 * 1024  # 20 MB
MAX_EXTRACTED_CHARACTERS = 500_000
MAX_PDF_PAGES = 500

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".csv", ".docx"}


class DocumentProcessingError(Exception):
    """Raised when a knowledge document cannot be processed."""


def _decode_text(raw_bytes: bytes) -> str:
    """Decode text with common encodings, without silently dropping bytes."""

    if not raw_bytes:
        return ""

    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise DocumentProcessingError(
        "Unable to decode this file as supported text."
    )


def _check_extracted_text(text: str) -> str:
    """Normalize extracted text and enforce a maximum character count."""

    text = text.replace("\x00", "").strip()

    if not text:
        raise DocumentProcessingError(
            "No readable text was found in this document."
        )

    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise DocumentProcessingError(
            "Extracted text exceeds the 500,000-character limit."
        )

    return text


def _extract_plain_text(file_obj) -> str:
    """Extract text from TXT, Markdown, or CSV content."""

    raw_bytes = file_obj.read()
    return _check_extracted_text(_decode_text(raw_bytes))


def _extract_pdf(file_obj) -> str:
    """Extract selectable text from a PDF document."""

    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentProcessingError(
            "PDF support is not installed. Install the pypdf dependency."
        ) from exc

    try:
        reader = PdfReader(file_obj, strict=True)

        if reader.is_encrypted:
            try:
                result = reader.decrypt("")
            except Exception as exc:
                raise DocumentProcessingError(
                    "This PDF is password-protected."
                ) from exc

            if not result:
                raise DocumentProcessingError(
                    "This PDF is password-protected."
                )

        if len(reader.pages) > MAX_PDF_PAGES:
            raise DocumentProcessingError(
                f"PDFs cannot exceed {MAX_PDF_PAGES} pages."
            )

        extracted_parts = []
        total_characters = 0

        for page in reader.pages:
            try:
                page_text = page.extract_text() or ""
            except Exception as exc:
                raise DocumentProcessingError(
                    "Unable to extract text from a PDF page."
                ) from exc

            total_characters += len(page_text)

            if total_characters > MAX_EXTRACTED_CHARACTERS:
                raise DocumentProcessingError(
                    "Extracted PDF text exceeds the 500,000-character limit."
                )

            if page_text.strip():
                extracted_parts.append(page_text.strip())

        return _check_extracted_text("\n\n".join(extracted_parts))

    except DocumentProcessingError:
        raise
    except Exception as exc:
        logger.warning("PDF extraction failed.")
        raise DocumentProcessingError(
            "The PDF is invalid or could not be read."
        ) from exc


def _extract_docx(file_obj) -> str:
    """Extract paragraph and table text from a DOCX document."""

    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentProcessingError(
            "DOCX support is not installed. Install the python-docx dependency."
        ) from exc

    try:
        document = Document(file_obj)
        parts = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip().replace("\n", " ")
                    for cell in row.cells
                ]
                if any(cells):
                    parts.append(" | ".join(cells))

        return _check_extracted_text("\n".join(parts))

    except DocumentProcessingError:
        raise
    except Exception as exc:
        logger.warning("DOCX extraction failed.")
        raise DocumentProcessingError(
            "The DOCX file is invalid or could not be read."
        ) from exc


def extract_uploaded_file(document: KnowledgeDocument) -> str:
    """
    Extract text from a stored, uploaded KnowledgeDocument.

    Supported extensions: PDF, TXT, Markdown, CSV, and DOCX.
    """

    if not document.file:
        raise DocumentProcessingError(
            "This document does not have an uploaded file."
        )

    filename = document.original_filename or document.file.name
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentProcessingError(
            "Unsupported file type. Supported: PDF, TXT, MD, CSV, DOCX."
        )

    try:
        document.file.open("rb")
        file_obj = document.file

        try:
            file_size = getattr(document, "size_bytes", None)

            if file_size is None:
                file_size = getattr(file_obj, "size", None)

            if file_size is not None and file_size > MAX_DOCUMENT_SIZE:
                raise DocumentProcessingError(
                    "Document size cannot exceed 20 MB."
                )

            if extension == ".pdf":
                return _extract_pdf(file_obj)

            if extension == ".docx":
                return _extract_docx(file_obj)

            return _extract_plain_text(file_obj)

        finally:
            file_obj.close()

    except DocumentProcessingError:
        raise
    except Exception as exc:
        logger.warning(
            "Could not read uploaded knowledge document %s.",
            document.pk,
        )
        raise DocumentProcessingError(
            "Unable to read the uploaded document."
        ) from exc


def extract_document_text(document: KnowledgeDocument) -> str:
    """
    Return source text for a KnowledgeDocument.

    Plain-text sources are handled directly. Uploaded files are extracted
    according to their extension. URL fetching is intentionally handled
    separately so a future fetcher can enforce SSRF protections.
    """

    if document.source_type == KnowledgeDocument.SourceType.TEXT:
        text = document.source_text or ""

        if len(text) > MAX_EXTRACTED_CHARACTERS:
            raise DocumentProcessingError(
                "Source text exceeds the 500,000-character limit."
            )

        return _check_extracted_text(text)

    if document.source_type == KnowledgeDocument.SourceType.UPLOAD:
        return extract_uploaded_file(document)

    if document.source_type == KnowledgeDocument.SourceType.URL:
        raise DocumentProcessingError(
            "URL processing is not implemented yet. "
            "A secure URL fetcher must be connected first."
        )

    raise DocumentProcessingError("Unsupported document source type.")

